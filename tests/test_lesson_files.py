"""
Tests for src/invoice_generator/lesson_files.py preparing a lesson's files.

The files are real images, PDFs and presentations made on the spot, so
the tests check what the encoders actually do, including when re-encoding
fails to make a file smaller.
"""

import logging
import random
import pytest
import typst
from PIL import Image
from pptx import Presentation

from invoice_generator.drive_api import DriveFile
from invoice_generator.lesson_files import (
    MAX_IMAGE_EDGE,
    Kind,
    Material,
    format_size,
    pdf_pages,
    prepare,
    shrink_image,
    slide_text,
    smallest,
)


def photo(width, height):
    """A noisy image, which compresses like a photo: badly."""
    rng = random.Random(0)
    return Image.frombytes(
        "RGB", (width, height), rng.randbytes(width * height * 3)
    )


def typst_pdf(path, pages):
    """A text-only PDF, like a typed handout, with `pages` pages."""
    source = path.with_suffix(".typ")
    source.write_text("#pagebreak()\n".join(["Integration by parts\n"] * pages))
    typst.compile(source, output=path)
    return path


@pytest.fixture
def workdir(tmp_path):
    directory = tmp_path / "work"
    directory.mkdir()
    return directory


class TestFormatSize:
    """Tests for format_size function."""

    @pytest.mark.parametrize("size, text", [
        (340_000, "340 kB"),
        (999_499, "999 kB"),
        (2_100_000, "2.1 MB"),
    ])
    def test_sizes(self, size, text):
        assert format_size(size) == text


class TestSmallest:
    """Tests for smallest function."""

    def test_picks_the_smallest_candidate(self, tmp_path, caplog):
        """A smaller re-encoding wins, and the saving is logged."""
        original, a, b = (tmp_path / name for name in ("o", "a", "b"))
        original.write_bytes(b"x" * 3_000_000)
        a.write_bytes(b"x" * 2_000_000)
        b.write_bytes(b"x" * 500_000)

        with caplog.at_level(logging.INFO):
            assert smallest(original, [a, b]) == b
        assert "Shrank 'o' from 3.0 MB to 500 kB" in caplog.text

    def test_keeps_the_original_if_nothing_is_smaller(self, tmp_path, caplog):
        """Re-encoding can make files bigger; then the original stays."""
        original, bigger = tmp_path / "o", tmp_path / "b"
        original.write_bytes(b"x" * 1000)
        bigger.write_bytes(b"x" * 2000)

        with caplog.at_level(logging.INFO):
            assert smallest(original, [bigger]) == original
        assert "Kept 'o' as it is (1 kB)" in caplog.text

    def test_keeps_the_original_in_a_tie(self, tmp_path):
        """No point swapping a file for one just as big."""
        original, same = tmp_path / "o", tmp_path / "s"
        original.write_bytes(b"x" * 1000)
        same.write_bytes(b"y" * 1000)

        assert smallest(original, [same]) == original


class TestShrinkImage:
    """Tests for shrink_image function."""

    def test_large_photos_are_scaled_down(self, tmp_path, workdir):
        """A big photo should come back smaller, and no larger than
        models look at."""
        original = tmp_path / "Photo of homework.png"
        photo(1800, 1200).save(original)

        shrunk = shrink_image(original, workdir)

        assert shrunk.stat().st_size < original.stat().st_size
        with Image.open(shrunk) as image:
            assert max(image.size) == MAX_IMAGE_EDGE

    def test_already_small_images_are_kept(self, tmp_path, workdir):
        """A small, well-compressed drawing can't usefully shrink: as a
        JPEG it would grow, so the original is kept."""
        original = tmp_path / "Diagram.png"
        Image.new("RGB", (200, 100), "white").save(original, optimize=True)

        assert shrink_image(original, workdir) == original

    def test_transparency_becomes_white(self, tmp_path, workdir):
        """Excalidraw exports have transparent backgrounds, which should
        come out white rather than black."""
        original = tmp_path / "Excalidraw.png"
        photo(1700, 1700).convert("RGBA").save(original)
        with Image.open(original) as image:
            image.putalpha(0)
            image.save(original)

        shrunk = shrink_image(original, workdir)

        assert shrunk != original
        with Image.open(shrunk) as image:
            assert image.convert("RGB").getpixel((10, 10)) == (255, 255, 255)

    def test_photos_are_turned_the_right_way_up(self, tmp_path, workdir):
        """Phones store some photos on their side, with an EXIF note of
        which way up they go; the shrunk version is the right way up."""
        original = tmp_path / "Sideways.jpg"
        exif = Image.Exif()
        exif[0x0112] = 6  # orientation: rotate 90° clockwise to view
        photo(1700, 900).save(original, exif=exif, quality=100)

        with Image.open(shrink_image(original, workdir)) as image:
            assert image.width < image.height


class TestPdfPages:
    """Tests for pdf_pages function."""

    def test_renders_the_first_pages(self, tmp_path, workdir, caplog):
        """Only the first `max_pages` pages are rendered, at no more than
        `max_edge` pixels, and the rest are noted in the log."""
        pdf = typst_pdf(tmp_path / "Handout.pdf", pages=3)

        with caplog.at_level(logging.INFO):
            pages = pdf_pages(pdf, workdir, max_pages=2, max_edge=800)

        assert len(pages) == 2
        for page in pages:
            with Image.open(page) as image:
                assert max(image.size) <= 800
        assert "first 2 of 3 pages of 'Handout.pdf'" in caplog.text

    def test_short_pdfs_are_rendered_whole(self, tmp_path, workdir):
        pdf = typst_pdf(tmp_path / "Handout.pdf", pages=1)

        assert len(pdf_pages(pdf, workdir, max_pages=5)) == 1


def presentation(path, *slides):
    """A PowerPoint file with a title-only slide per entry in `slides`
    (None for a slide with no text)."""
    deck = Presentation()
    for title in slides:
        slide = deck.slides.add_slide(deck.slide_layouts[5])
        if title is not None:
            slide.shapes.title.text = title
    deck.save(str(path))
    return path


class TestSlideText:
    """Tests for slide_text function."""

    def test_text_is_kept_slide_by_slide(self, tmp_path, workdir):
        """Slides with text are numbered, slides without are skipped."""
        deck = presentation(tmp_path / "Vectors.pptx", "Vectors", None, "Dot product")

        text = slide_text(deck, workdir)

        assert text is not None
        assert text.read_text() == "Slide 1: Vectors\nSlide 3: Dot product"

    def test_presentations_without_text_give_nothing(self, tmp_path, workdir):
        deck = presentation(tmp_path / "Pictures.pptx", None)

        assert slide_text(deck, workdir) is None


class TestPrepare:
    """Tests for prepare function."""

    def test_each_kind_of_file(self, tmp_path, workdir):
        """Every file comes back in order, under its Drive name, as the
        kind of thing the model can read; PDFs as they are, as they're
        rendered as page images later."""
        pdf = typst_pdf(tmp_path / "Handout.pdf", pages=1)
        png = tmp_path / "Diagram.png"
        Image.new("RGB", (200, 100), "white").save(png, optimize=True)
        txt = tmp_path / "README.txt"
        txt.write_text("Read me")
        pptx = presentation(tmp_path / "Vectors.pptx", "Vectors")
        heic = tmp_path / "Photo.heic"
        heic.write_bytes(b"not something we can read")

        materials = prepare([
            DriveFile("Handout.pdf", pdf),
            DriveFile("Extra/Diagram.png", png),
            DriveFile("README", txt),
            DriveFile("Vectors.pptx", pptx),
            DriveFile("Photo.heic", heic),
            DriveFile("Quiz", None),
        ], workdir)

        assert [(m.name, m.kind) for m in materials] == [
            ("Handout.pdf", Kind.PDF),
            ("Extra/Diagram.png", Kind.IMAGE),
            ("README", Kind.TEXT),
            ("Vectors.pptx", Kind.TEXT),
            ("Photo.heic", Kind.OTHER),
            ("Quiz", Kind.OTHER),
        ]
        assert materials[0].path == pdf
        assert materials[2].path == txt
        assert materials[3].path is not None
        assert materials[3].path.read_text() == "Slide 1: Vectors"
        assert materials[4].path is None

    def test_presentations_without_text_go_by_name(self, tmp_path, workdir):
        deck = presentation(tmp_path / "Pictures.pptx", None)

        assert prepare([DriveFile("Pictures.pptx", deck)], workdir) == [
            Material("Pictures.pptx", Kind.OTHER, None)
        ]

    def test_files_that_only_claim_a_format_go_by_name(self, tmp_path, workdir, caplog):
        """A damaged image is logged, and summarised by its name."""
        fake = tmp_path / "Photo.png"
        fake.write_bytes(b"not an image")

        with caplog.at_level(logging.WARNING):
            materials = prepare([DriveFile("Photo.png", fake)], workdir)

        assert materials == [Material("Photo.png", Kind.OTHER, None)]
        assert "Could not read 'Photo.png'" in caplog.text

    def test_files_with_the_same_name_do_not_clash(self, tmp_path, workdir):
        """Two "Photo.png"s, from different subfolders, stay distinct."""
        (tmp_path / "a").mkdir()
        (tmp_path / "b").mkdir()
        photo(1700, 900).save(tmp_path / "a" / "Photo.png")
        photo(900, 1700).save(tmp_path / "b" / "Photo.png")

        first, second = prepare([
            DriveFile("A/Photo.png", tmp_path / "a" / "Photo.png"),
            DriveFile("B/Photo.png", tmp_path / "b" / "Photo.png"),
        ], workdir)

        assert first.path != second.path
