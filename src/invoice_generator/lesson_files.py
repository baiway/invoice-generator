"""
Preparing a lesson's files to be summarised.

Apple's on-device model reads images and text, so PDFs are rendered as page
images (`pdf_pages`), large images are scaled down to what the model can
use, and PowerPoint files are reduced to their text. Re-encoding doesn't
always make an image smaller (an already well-compressed JPEG can grow, as
can a line drawing saved as a JPEG), so the smaller version is only used if
it really is smaller.
"""

from dataclasses import dataclass
from enum import Enum
from pathlib import Path

import pypdfium2
from PIL import Image, ImageOps
from pptx import Presentation

from invoice_generator.drive_api import DriveFile
from invoice_generator.logging_config import get_logger

logger = get_logger(__name__)

# Images are scaled down to at most this many pixels on the long edge:
# plenty for the model to read, and far smaller than a phone's photos
MAX_IMAGE_EDGE = 1568
JPEG_QUALITY = 80

IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".gif", ".webp"}
TEXT_SUFFIXES = {".txt", ".md", ".csv", ".tex"}


class Kind(Enum):
    """What can be done with a file to summarise it."""

    PDF = "pdf"
    IMAGE = "image"
    TEXT = "text"
    # Nothing but its name: a format the model can't read, or a file that
    # couldn't be downloaded or opened
    OTHER = "other"


@dataclass(frozen=True)
class Material:
    """One of a lesson's files, ready to summarise."""

    # As in Drive, which is often the best guide to what the lesson covered
    name: str
    kind: Kind
    # The file to read, or None for `Kind.OTHER`
    path: Path | None


def format_size(size: int) -> str:
    """A file size for the log, e.g. "340 kB" or "2.1 MB"."""
    if size < 1_000_000:
        return f"{size / 1000:.0f} kB"
    return f"{size / 1_000_000:.1f} MB"


def smallest(original: Path, candidates: list[Path]) -> Path:
    """The smallest of `original` and the re-encoded `candidates`,
    preferring `original` in a tie. Logs what was saved, if anything."""
    best = min([original, *candidates], key=lambda path: path.stat().st_size)
    before = format_size(original.stat().st_size)
    if best == original:
        logger.info(f"Kept '{original.name}' as it is ({before}); "
                    "re-encoding made nothing smaller")
    else:
        logger.info(f"Shrank '{original.name}' from {before} to "
                    f"{format_size(best.stat().st_size)}")
    return best


def flatten(image: Image.Image) -> Image.Image:
    """An RGB copy of `image`, transparent areas made white (Excalidraw
    exports, for one, have transparent backgrounds)."""
    image = image.convert("RGBA")
    background = Image.new("RGBA", image.size, "white")
    return Image.alpha_composite(background, image).convert("RGB")


def shrink_image(path: Path, workdir: Path, max_edge: int = MAX_IMAGE_EDGE) -> Path:
    """Scale an image down to `max_edge` pixels on its long edge, and save
    it as both a PNG (best for drawings and screenshots) and a JPEG (best
    for photos).

    Args:
        path: The image
        workdir: Directory to write the smaller versions to
        max_edge: Longest edge, in pixels, worth keeping

    Returns:
        Whichever of the image and its smaller versions is smallest
    """
    with Image.open(path) as opened:
        # Photos are often stored on their side, with a note of which way
        # up they go; the re-encoded versions are the right way up
        image = flatten(ImageOps.exif_transpose(opened))
    image.thumbnail((max_edge, max_edge))

    png = workdir / f"{path.stem} (shrunk).png"
    jpeg = workdir / f"{path.stem} (shrunk).jpg"
    image.save(png, optimize=True)
    image.save(jpeg, quality=JPEG_QUALITY, optimize=True)
    return smallest(path, [png, jpeg])


def pdf_pages(
    path: Path,
    workdir: Path,
    max_pages: int,
    max_edge: int = MAX_IMAGE_EDGE,
) -> list[Path]:
    """Render the first `max_pages` pages of a PDF as images, for the
    model to read.

    Args:
        path: The PDF
        workdir: Directory to write the page images to
        max_pages: How many pages to render, from the start
        max_edge: Longest edge of each page image, in pixels

    Returns:
        The page images, in order, each as small as `shrink_image` can
        make it
    """
    pages = []
    pdf = pypdfium2.PdfDocument(path)
    try:
        if len(pdf) > max_pages:
            logger.info(f"Reading only the first {max_pages} of "
                        f"{len(pdf)} pages of '{path.name}'")
        for number in range(min(len(pdf), max_pages)):
            page = pdf[number]
            scale = max_edge / max(page.get_size())
            rendered = workdir / f"{path.stem} page {number + 1}.png"
            page.render(scale=scale).to_pil().save(rendered)
            pages.append(shrink_image(rendered, workdir, max_edge))
    finally:
        pdf.close()
    return pages


def slide_text(path: Path, workdir: Path) -> Path | None:
    """The text of a PowerPoint presentation, slide by slide, saved as a
    text file, or None if it has no text."""
    lines = []
    for number, slide in enumerate(Presentation(str(path)).slides, start=1):
        texts = [
            shape.text_frame.text.strip()
            for shape in slide.shapes
            if shape.has_text_frame and shape.text_frame.text.strip()
        ]
        if texts:
            lines.append(f"Slide {number}: " + " / ".join(texts))
    if not lines:
        return None
    text = workdir / f"{path.stem} (slides).txt"
    text.write_text("\n".join(lines))
    return text


def prepare(files: list[DriveFile], workdir: Path) -> list[Material]:
    """Prepare a lesson's files to be summarised.

    Args:
        files: The lesson's files, from `drive_api.download_folder`
        workdir: Directory to write smaller versions and extracted text to

    Returns:
        Each file, in the same order, made as small as it usefully can be
    """
    materials = []
    for number, file in enumerate(files):
        material = Material(file.name, Kind.OTHER, None)
        path = file.path
        suffix = path.suffix.lower() if path else ""
        # A directory per file, as files in different subfolders can share
        # a name
        filedir = workdir / str(number)
        filedir.mkdir(parents=True, exist_ok=True)
        try:
            if path is None:
                pass
            elif suffix == ".pdf":
                # Rendered as page images when summarised, so there's no
                # point compressing them first
                material = Material(file.name, Kind.PDF, path)
            elif suffix in IMAGE_SUFFIXES:
                material = Material(file.name, Kind.IMAGE, shrink_image(path, filedir))
            elif suffix in TEXT_SUFFIXES:
                material = Material(file.name, Kind.TEXT, path)
            elif suffix == ".pptx":
                text = slide_text(path, filedir)
                if text:
                    material = Material(file.name, Kind.TEXT, text)
        except Exception as e:
            # Damaged files, or ones that only claim to be a format. The
            # name can still be summarised.
            logger.warning(f"Could not read '{file.name}', so going by its name: {e}")
        if material.kind is Kind.OTHER:
            logger.info(f"Going by the name alone of '{file.name}'")
        materials.append(material)
    return materials
